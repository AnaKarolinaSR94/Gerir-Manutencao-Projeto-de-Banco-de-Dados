# -*- coding: utf-8 -*-
"""Parte 3 (Projeto Fisico em PostgreSQL) do documento do projeto.

Monta a Parte 3 com os mesmos estilos do build_docx.py e com os resultados reais das consultas,
obtidos executando os scripts da pasta sql/ em um banco PostgreSQL local.

Pre-requisito: PostgreSQL rodando e um banco vazio ou descartavel (padrao: gerir_manutencao;
use as variaveis PGDATABASE, PGHOST, PGUSER... para outro). O schema sgm e recriado do zero.

Uso 1 - acrescentar a Parte 3 a um documento ja existente (antes das Referencias):
    python ferramentas/parte3_docx.py entregas/Projeto_BD_GM_Gerir_Manutencao.docx [saida.docx]
Uso 2 - dentro do build_docx.py:  import parte3_docx; parte3_docx.add_parte3(doc)
"""
import copy
import datetime
import os
import re
import sys

import psycopg2
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from schema import TABLES, fks  # noqa: E402
from gen_ddl import EXTRA_CHECKS, INDEXES, on_delete, ordem_topologica  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LARGURA = 16.0      # cm uteis da pagina retrato (A4 com margens de 2,5 cm)
MAX_LINHAS = 15     # linhas exibidas por resultado; o restante fica no arquivo sql/resultados/
HEADER_FILL = "D9E2F3"
CODE_FILL = "F2F2F2"


# ---------------------------------------------------------------- helpers (mesmo visual do build_docx.py)
class Escritor:
    def __init__(self, doc):
        self.doc = doc

    def H(self, text, lvl):
        return self.doc.add_heading(text, lvl)

    def P(self, text="", bold=False, italic=False, size=None, align=None, space_after=6, keep=False):
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = keep
        return p

    def PJ(self, text):
        return self.P(text, align="justify")

    def B(self, text, bold_prefix=None):
        p = self.doc.add_paragraph(style="List Bullet")
        if bold_prefix:
            p.add_run(bold_prefix).bold = True
        p.add_run(text)
        p.paragraph_format.space_after = Pt(2)
        return p

    def CODE(self, text, size=8, keep=False):
        p = self.doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")  # sombreado antes da formatacao: o python-docx encaixa o resto na ordem do schema
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), CODE_FILL)
        ppr.append(shd)
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(2), Pt(6)
        pf.left_indent = Cm(0.2)
        pf.line_spacing = 1.0
        pf.keep_with_next = keep
        lines = text.split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(size)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            if i < len(lines) - 1:
                r.add_break()
        return p

    def T(self, headers, rows, widths=None, font=9):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(str(h))
            r.bold = True
            r.font.size = Pt(font)
            shade(c, HEADER_FILL)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(font)
        _repeat_header(t)
        if widths:
            for i, w in enumerate(widths):
                t.columns[i].width = Cm(w)
                for row in t.rows:
                    row.cells[i].width = Cm(w)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
        return t

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def margens_estreitas(t, cm=0.07):
    """margem interna das celulas de 0,07 cm (o padrao do Word e 0,19 cm), para caber mais colunas"""
    tblPr = t._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for lado in ("left", "right"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(int(cm * 567))); el.set(qn("w:type"), "dxa")
        mar.append(el)
    look = tblPr.find(qn("w:tblLook"))  # ordem do schema: ... tblLayout, tblCellMar, tblLook
    if look is not None:
        look.addprevious(mar)
    else:
        tblPr.append(mar)


def _repeat_header(t):
    trPr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "t" if v else "f"
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    return str(v)


def larguras(headers, rows, total=LARGURA, font=7, pad=0.3):
    """largura de cada coluna proporcional ao conteudo, sem quebrar palavras curtas"""
    cw = font * 0.0215  # cm por caractere (Calibri, com folga para maiusculas e digitos)
    nat, mn = [], []
    for i, h in enumerate(headers):
        vals = [str(r[i]) for r in rows]
        m = max([len(h)] + [len(v) for v in vals])
        palavra = max([len(w) for v in vals for w in v.split()] or [1])
        nat.append(m * cw + pad)
        mn.append(min(max(palavra, min(len(h), 9), 3), 18) * cw + pad)
    if sum(nat) <= total:
        return nat
    if sum(mn) >= total:
        return [m * total / sum(mn) for m in mn]
    sobra, extra = total - sum(mn), [n - m for n, m in zip(nat, mn)]
    return [m + sobra * e / sum(extra) for m, e in zip(mn, extra)]


def ler(rel):
    return open(os.path.join(ROOT, rel), encoding="utf-8").read()


def split_statements(sql):
    stmts, cur, dollar = [], [], False
    for line in sql.split("\n"):
        cur.append(line)
        if line.count("$$") % 2:
            dollar = not dollar
        if not dollar and re.sub(r"\s+--.*$", "", line).rstrip().endswith(";"):
            stmts.append("\n".join(cur)); cur = []
    if "\n".join(cur).strip():
        stmts.append("\n".join(cur))
    return [s for s in stmts if any(l.strip() and not l.strip().startswith("--") for l in s.split("\n"))]


# ---------------------------------------------------------------- banco
class Banco:
    def __init__(self):
        self.conn = psycopg2.connect(dbname=os.environ.get("PGDATABASE", "gerir_manutencao"))
        self.conn.autocommit = True
        self.cur = self.conn.cursor()

    def run(self, stmt):
        self.conn.notices.clear()
        self.cur.execute(stmt)
        cols = [d[0] for d in self.cur.description] if self.cur.description else None
        rows = self.cur.fetchall() if cols else None
        notices = [n.strip().replace("NOTICE:  ", "") for n in self.conn.notices]
        return cols, rows, self.cur.statusmessage, notices


# ---------------------------------------------------------------- blocos de script executados
def bloco_resultado(w, cols, rows):
    shown = rows[:MAX_LINHAS]
    data = [[fmt(v) for v in r] for r in shown]
    font = 7
    while font > 5.5:  # tabela muito larga usa fonte menor em vez de quebrar codigos e datas
        cw = font * 0.0215
        minimo = sum(min(max(max([len(x) for v in [str(r[i]) for r in data] for x in v.split()] or [1]),
                             min(len(c), 9), 3), 18) * cw + 0.3 for i, c in enumerate(cols))
        if minimo <= LARGURA:
            break
        font -= 0.5
    t = w.T(cols, data, widths=larguras(cols, data, font=font), font=font)
    margens_estreitas(t)
    n = len(rows)
    nota = f"({n} {'linha' if n == 1 else 'linhas'}"
    nota += f"; exibidas as {MAX_LINHAS} primeiras – resultado completo em sql/resultados/)" if n > MAX_LINHAS else ")"
    w.P(nota, italic=True, size=8, space_after=8)


def titulo(cap, rotulo=""):
    t = " ".join(cap)
    m = re.match(r"^(RF\d+|\d+\.\d+(?: \(b\))?)\s*-?\s*(.*)$", t)
    if m:
        num, resto = m.groups()
        pre = "" if num.startswith("RF") else rotulo  # "Consulta 1.1" / "Etapa 3.4" nao se confundem com as secoes
        return f"{pre}{num} – {resto[:1].upper() + resto[1:]}"
    return t[:1].upper() + t[1:]


def secao_executada(w, db, sql, sub_prefixo, sub_inicial=1, tabela_do=None, rotulo=""):
    """Executa o script bloco a bloco (blocos separados por linha em branco) e escreve:
    titulo (comentario do bloco), SQL e resultado. Blocos so de comentario viram subtitulos (H3)."""
    n_sub = sub_inicial - 1
    for k, par in enumerate(re.split(r"\n\s*\n", sql)):
        lines = par.strip("\n").split("\n")
        cap = []
        while lines and lines[0].strip().startswith("--"):
            cap.append(lines.pop(0).strip()[2:].strip())
        body = "\n".join(lines).strip()
        cap = [c for c in cap if c and not re.fullmatch(r"[=\-]+", c)]
        if k == 0 and not body:
            continue  # cabecalho do arquivo
        if cap and re.match(r"^(PARTE \d|\d\. )", cap[0]):  # comentario de parte -> subtitulo (H3)
            n_sub += 1
            nome = re.sub(r"^(PARTE \d+ - |\d+\. )", "", cap[0])
            w.H(f"{sub_prefixo}.{n_sub} {nome[:1].upper() + nome[1:].lower() if nome.isupper() else nome}", 3)
            for c in cap[1:]:
                w.P(c[:1].upper() + c[1:], italic=True, size=10, keep=True)
            cap = []
        if not body:
            continue
        if body.upper().startswith("SET SEARCH_PATH") or body.upper() in ("BEGIN;", "ROLLBACK;"):
            db.run(body)
            continue
        if body.startswith("DO $$") and tabela_do:
            tabela_do(w, db, cap, body)
            continue
        if cap:
            w.P(titulo(cap, rotulo), bold=True, size=10, space_after=2, keep=True)
        stmts = split_statements(body)
        w.CODE(body, keep=True)
        status = []
        for st in stmts:
            cols, rows, msg, notices = db.run(st)
            if cols is not None:
                if status:
                    w.P("Resultado: " + "; ".join(status), italic=True, size=8, space_after=4, keep=True); status = []
                bloco_resultado(w, cols, rows)
            elif msg and not msg.startswith(("SET", "DO")):
                status.append(msg)
        if status:
            w.P("Resultado: " + "; ".join(status), italic=True, size=8, space_after=8)


def tabela_restricoes(w, db, cap, body):
    """bloco DO com as tentativas invalidas: mostra em tabela (tentativa x mensagem do PostgreSQL)"""
    tentativas = re.findall(r"BEGIN\s+--\s*(.+)", body)
    _, _, _, notices = db.run(body)
    w.P("O bloco DO do script tenta executar sete operações inválidas. Cada tentativa fica em um sub-bloco "
         "com tratamento de exceção: se a restrição funcionar, o erro é capturado e registrado como NOTICE, sem "
         "abortar a transação. Resultado da execução:", align="justify", keep=True)
    rows = []
    for t, n in zip(tentativas, notices):
        ok = n.startswith("OK")
        msg = re.sub(r"^(OK|ERRO): [^(]*\(", "", n).rstrip(")")
        rows.append([t[:1].upper() + t[1:], "Bloqueada" if ok else "ACEITA (falha)", msg])
    w.T(["Tentativa", "Resultado", "Mensagem do PostgreSQL"], rows, widths=[4.4, 2.0, 9.6], font=8)
    w.P("", space_after=6)


# ---------------------------------------------------------------- conteudo da Parte 3
def add_parte3(doc):
    w = Escritor(doc)
    db = Banco()
    ddl, carga = ler("sql/01_criacao.sql"), ler("sql/02_insercao.sql")
    db.run(ddl)
    db.run(carga)

    w.H("PARTE 3 – PROJETO FÍSICO", 1)
    w.PJ("Esta parte apresenta o projeto físico do banco de dados do Gerir Manutenção (SGM) no SGBD PostgreSQL 16, "
         "obtido a partir do esquema relacional normalizado da Parte 2. São apresentados as decisões de "
         "implementação, o script de criação das tabelas, a carga de dados de teste, as consultas (com os "
         "resultados obtidos na execução) e as operações de manipulação de dados que reproduzem o fluxo do sistema.")

    # 3.1 ------------------------------------------------------------------
    w.H("3.1 Ambiente e organização dos scripts", 2)
    w.PJ("Todas as tabelas são criadas no schema sgm. O script de criação é gerado automaticamente a partir de "
         "ferramentas/schema.py, a mesma fonte usada para o DER relacional (seção 2.2) e para o dicionário de dados "
         "(seção 2.3); assim, o banco implementado corresponde exatamente ao modelo lógico documentado. Os scripts "
         "ficam na pasta sql/ do repositório:")
    w.T(["Arquivo", "Conteúdo"], [
        ["01_criacao.sql", "Criação do schema, das 19 tabelas (chaves, restrições, valores padrão) e dos índices"],
        ["02_insercao.sql", "Carga de dados de teste, coerente com as regras de negócio"],
        ["03_consultas.sql", "Consultas de exemplo, consultas de exercício e consultas dos requisitos funcionais"],
        ["04_manipulacao.sql", "Inserções, atualizações e exclusões no fluxo do sistema e testes das restrições"],
        ["00_executar_tudo.sql", "Executa os quatro scripts em ordem pelo psql"],
        ["resultados/", "Saída completa da execução de 03 e 04"],
    ], widths=[4.0, 12.0], font=9)
    w.P("", space_after=2)
    w.P("Execução pelo terminal (a partir da raiz do repositório):", space_after=2, keep=True)
    w.CODE("createdb gerir_manutencao\npsql -d gerir_manutencao -f sql/00_executar_tudo.sql", size=9)
    w.PJ("Em ferramentas gráficas (pgAdmin, DBeaver), basta criar o banco gerir_manutencao e executar os "
         "arquivos 01, 02, 03 e 04 nessa ordem.")

    # 3.2 ------------------------------------------------------------------
    w.H("3.2 Decisões do projeto físico", 2)
    w.H("3.2.1 Tipos de dados e chaves", 3)
    for pre, txt in [
        ("Chaves substitutas: ", "as entidades fortes usam SERIAL (inteiro com sequência automática) como chave primária."),
        ("Especializações: ", "as subclasses (TECNICO, SUPERVISOR, GESTOR, ADMINISTRATIVO, ADMINISTRADOR_SISTEMA, "
                              "MANUTENCAO_CORRETIVA e MANUTENCAO_PREVENTIVA) usam a chave da superclasse como chave "
                              "primária e estrangeira ao mesmo tempo."),
        ("Datas e horas: ", "TIMESTAMP para eventos que entram no cálculo de MTBF, MTTR e horas trabalhadas "
                            "(falha, início e fim da execução, avaliação, baixa); DATE para prazos e vencimentos."),
        ("Domínios fechados: ", "listas como criticidade, prioridade, status e periodicidade são VARCHAR/CHAR com "
                                "CHECK ... IN (...), que deixam os valores válidos visíveis no próprio DDL."),
        ("Medidas: ", "NUMERIC(10,1) para o horímetro (sem erro de arredondamento) e INTEGER com CHECK ≥ 0 para "
                      "as quantidades de estoque."),
        ("Segurança: ", "a senha é armazenada apenas como hash (senha_hash), conforme o requisito de LGPD."),
    ]:
        w.B(txt, pre)

    w.H("3.2.2 Ações referenciais (ON DELETE)", 3)
    w.PJ("Cada chave estrangeira recebeu uma ação de exclusão de acordo com o papel do relacionamento: chave "
         "estrangeira que faz parte da chave primária (herança, entidade fraca e relacionamentos N:N) usa CASCADE; "
         "chave estrangeira opcional usa SET NULL; as demais usam RESTRICT. Duas exceções preservam o histórico de "
         "manutenção: a baixa de peça e o resultado de checklist não podem perder a referência à peça e ao item "
         "de checklist. Na prática, equipamentos, peças e usuários com histórico não são excluídos, e sim "
         "desativados (atributo ativo).")
    rows = []
    for t in ordem_topologica(TABLES):
        for local, ref_t, ref_c in fks(t):
            rows.append([t["name"], ", ".join(local), f"{ref_t}({', '.join(ref_c)})", on_delete(t, local)])
    w.T(["Tabela", "Chave estrangeira", "Referencia", "ON DELETE"], rows, widths=[4.2, 4.6, 5.0, 2.2], font=8)
    w.P("", space_after=2)

    w.H("3.2.3 Restrições de integridade adicionais", 3)
    w.PJ("Além das restrições já descritas no dicionário de dados (seção 2.3), o projeto físico acrescenta "
         "restrições CHECK que garantem regras de negócio do fluxo da ordem de serviço:")
    regra = {
        "data_fim_execucao": "O fim da execução não pode ser anterior ao início",
        "(status": "A OS só fica CONCLUIDA quando a decisão vigente é APROVADA, e vice-versa (RF12)",
        "decisao_avaliacao IS NULL": "Toda decisão registra a data e o supervisor que avaliou",
        "data_leitura": "Uma notificação só pode ser lida depois de enviada",
    }
    rows = []
    for tab, cks in EXTRA_CHECKS.items():
        for ck in cks:
            desc = next(v for k, v in regra.items() if ck.startswith(k))
            rows.append([tab, ck, desc])
    w.T(["Tabela", "Restrição (CHECK)", "Regra de negócio"], rows, widths=[3.2, 7.6, 5.2], font=8)
    w.P("", space_after=2)
    w.PJ("Destaque para a segunda regra: a forma direta status <> 'CONCLUIDA' OR decisao_avaliacao = 'APROVADA' "
         "deixaria passar uma OS concluída sem decisão, porque a comparação com NULL resulta em desconhecido e o "
         "CHECK aceita valores desconhecidos. Por isso a regra usa coalesce. O teste da seção 3.6 confirma o bloqueio.")

    w.H("3.2.4 Índices", 3)
    w.PJ("O PostgreSQL já cria índices para as chaves primárias e para as colunas UNIQUE. Foram criados índices "
         "adicionais nas colunas usadas pelas consultas mais frequentes do sistema:")
    uso = {
        "idx_os_equipamento": "Histórico da máquina (RF08), MTBF/MTTR e ranking (RF05, RF15)",
        "idx_os_tecnico": "Tarefas e produtividade por técnico (RF09, RF16), status da equipe (RF14)",
        "idx_os_status": "OS aguardando aprovação (RF12) e disponibilidade das máquinas (RF05)",
        "idx_plano_proxima_data": "Tarefas da semana (RF09) e alertas de proximidade/atraso (RF06)",
        "idx_baixa_peca": "Relatório de consumo de peças por período (RF07)",
        "idx_notificacao_data": "Supressão de alertas repetidos em 24 h (RF13)",
        "idx_tecnico_equipe": "Painel de status da equipe (RF14) e produtividade por equipe (RF16)",
    }
    w.T(["Índice", "Tabela", "Colunas", "Consultas atendidas"],
        [[n, t, c, uso[n]] for n, t, c in INDEXES], widths=[3.6, 3.0, 3.6, 5.8], font=8)
    w.P("", space_after=2)

    # 3.3 ------------------------------------------------------------------
    w.H("3.3 Script de criação do banco de dados", 2)
    w.PJ("Script completo (sql/01_criacao.sql). O comando CREATE DATABASE é executado uma única vez, conectado ao "
         "banco postgres; o restante é executado já conectado ao banco gerir_manutencao.")
    corpo = ddl.split("DROP SCHEMA", 1)[1]
    w.CODE("CREATE DATABASE gerir_manutencao;\n\nDROP SCHEMA" + corpo.rstrip(), size=7.5)

    # 3.4 ------------------------------------------------------------------
    w.H("3.4 Script de inserção de dados", 2)
    n_ins = carga.count("\nINSERT INTO")
    w.PJ(f"A carga de teste (sql/02_insercao.sql, {n_ins} comandos INSERT) é gerada por ferramentas/gen_carga.py, "
         "com dados fictícios e semente aleatória fixa, de modo que o resultado é sempre o mesmo. Cobre o período de "
         "abril a outubro de 2026 e usa 19/10/2026 08:00 como data de referência (“hoje”). Os dados respeitam as "
         "regras de negócio:")
    for txt in [
        "toda OS CONCLUIDA foi APROVADA; as OS rejeitadas voltaram para EM_EXECUCAO com justificativa;",
        "técnico com OS em execução está OCUPADO e os demais estão DISPONIVEL;",
        "o horímetro de cada equipamento cresce a cada manutenção, e horimetro_atual é a última leitura;",
        "o saldo de cada peça é o estoque inicial menos a soma das baixas, e o alerta de estoque mínimo foi gerado "
        "na baixa em que a peça atingiu o mínimo;",
        "cada preventiva executada tem o checklist do seu próprio equipamento respondido;",
        "há casos para todas as consultas: OS em todos os status, preventivas atrasadas, peças abaixo do mínimo, "
        "um técnico sem nenhuma OS e peças nunca utilizadas.",
    ]:
        w.B(txt)
    rows = []
    for t in TABLES:
        _, r, _, _ = db.run(f"SELECT count(*) FROM sgm.{t['name']}")
        rows.append([t["name"], r[0][0]])
    meio = (len(rows) + 1) // 2
    pares = [a + (b if b else ["", ""]) for a, b in zip(rows[:meio], rows[meio:] + [None])]
    w.P("Quantidade de registros por tabela após a carga:", space_after=2, keep=True)
    w.T(["Tabela", "Registros", "Tabela", "Registros"], pares, widths=[4.6, 2.0, 4.6, 2.0], font=8)
    w.P("", space_after=2)
    w.PJ("Trecho do script, com os primeiros comandos de algumas tabelas (o arquivo completo acompanha a entrega):")
    w.CODE(trecho_carga(carga), size=7.5)

    # 3.5 ------------------------------------------------------------------
    w.H("3.5 Consultas", 2)
    w.PJ("Consultas do arquivo sql/03_consultas.sql, executadas sobre a carga de dados da seção 3.4. As consultas que "
         "dependem da data atual usam a data de referência da carga (19/10/2026 08:00), para que os resultados sejam "
         "reproduzíveis; no sistema, esse valor é substituído por now() ou CURRENT_DATE. Resultados com mais de "
         f"{MAX_LINHAS} linhas aparecem resumidos; a saída completa está em sql/resultados/.")
    secao_executada(w, db, ler("sql/03_consultas.sql"), "3.5", rotulo="Consulta ")

    # 3.6 ------------------------------------------------------------------
    w.H("3.6 Manipulação de dados", 2)
    w.PJ("O arquivo sql/04_manipulacao.sql executa inserções, atualizações e exclusões na ordem em que acontecem no "
         "sistema: cadastros, o ciclo de vida de uma OS corretiva (distribuição, execução, baixa de peça com alerta de "
         "estoque mínimo, rejeição e aprovação) e de uma OS preventiva (checklist, aprovação e reagendamento pelo "
         "plano), além de testes das restrições e exclusões. Todo o script roda em uma transação terminada em "
         "ROLLBACK, para que a base volte ao estado da carga; para gravar as alterações, troca-se ROLLBACK por COMMIT.")
    secao_executada(w, db, ler("sql/04_manipulacao.sql"), "3.6", tabela_do=tabela_restricoes, rotulo="Etapa ")
    db.conn.close()


def trecho_carga(carga):
    escolha = {"usuario": 3, "equipe": 3, "equipamento": 2, "plano_preventivo": 2, "ordem_servico": 3,
               "manutencao_corretiva": 2, "baixa_peca": 2, "notificacao": 2}
    out = ["SET search_path TO sgm;", ""]
    for tab, n in escolha.items():
        linhas = [l for l in carga.split("\n") if l.startswith(f"INSERT INTO {tab} (")]
        out += linhas[:n] + ["..."]
    return "\n".join(out)


# ---------------------------------------------------------------- acrescentar a um .docx existente
def acrescentar(entrada, saida):
    doc = Document(entrada)
    body = doc.element.body
    paras = doc.paragraphs

    # capa: inclui a Entrega 3 depois da Entrega 2
    if not any(p.text.strip().startswith("Entrega 3") for p in paras[:40]):
        e2 = next(p for p in paras[:40] if p.text.strip().startswith("Entrega 2"))
        novo = copy.deepcopy(e2._p)
        e2._p.addnext(novo)
        e2.paragraph_format.space_after = Pt(0)
        e3 = next(p for p in doc.paragraphs[:40] if p._p is novo)
        for r in e3.runs[1:]:
            r._r.getparent().remove(r._r)
        e3.runs[0].text = "Entrega 3 – Projeto Físico"

    ref = next(p for p in doc.paragraphs if p.text.strip() == "Referências" and p.style.name.startswith("Heading"))
    antes = len(body)
    w = Escritor(doc)
    w.page_break()
    add_parte3(doc)
    w.page_break()
    # move o conteudo novo (adicionado no fim do corpo, antes do sectPr) para antes das Referencias
    novos = list(body)[antes - 1:-1] if body[-1].tag == qn("w:sectPr") else list(body)[antes:]
    novos = [e for e in novos if e.tag != qn("w:sectPr")]
    for e in novos:
        ref._p.addprevious(e)

    # pede ao Word para atualizar o sumario ao abrir
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        # ordem exigida pelo schema: updateFields vem antes de hdrShapeDefaults/footnotePr/.../compat/rsids...
        depois = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                  "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                  "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                  "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]
        alvo = next((e for e in settings if e.tag in [qn("w:" + d) for d in depois]), None)
        if alvo is not None:
            alvo.addprevious(uf)
        else:
            settings.append(uf)
    doc.save(saida)
    print("Parte 3 acrescentada:", saida)


if __name__ == "__main__":
    entrada = sys.argv[1] if len(sys.argv) > 1 else "entregas/Projeto_BD_GM_Gerir_Manutencao.docx"
    acrescentar(entrada, sys.argv[2] if len(sys.argv) > 2 else entrada)
