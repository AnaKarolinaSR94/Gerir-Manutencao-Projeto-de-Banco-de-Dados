# -*- coding: utf-8 -*-
"""Gera o documento Word das Entregas 1 (Projeto Logico Conceitual) e 2 (Projeto Logico Relacional).
Revisao de 11/09/2026: escopo alinhado aos documentos do grupo (requisitos/*.md) e modelo simplificado (ADR 0004)."""
import sys
sys.path.insert(0, "ferramentas")
import pymupdf
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from schema import TABLES, formal, fks, is_pk, fk_target, flags

OUT = "entregas/Projeto_BD_GM_Gerir_Manutencao.docx"
SIS = "Gerir Manutenção (SGM)"
doc = Document()

# ---------------------------------------------------------------- estilos base
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for lvl, size in ((1, 18), (2, 14), (3, 12)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.5)


# ---------------------------------------------------------------- helpers
def H(text, lvl=1):
    return doc.add_heading(text, lvl)


def P(text="", bold=False, italic=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    return p


def PJ(text):
    return P(text, align="justify")


def B(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)
    return p


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def T(headers, rows, widths=None, font=9, header_fill="D9E2F3"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        shade(c, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            rr = cells[i].paragraphs[0].add_run(str(v))
            rr.font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def IMG(path, width_cm, caption=None):
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        P(caption, italic=True, size=9, align="center", space_after=10)


def landscape(on=True):
    s = doc.add_section()
    if on:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(2.0)
    return s


def FORMAL(table, font=10):
    """Paragrafo com a relacao em notacao formal: PK sublinhada, FK com #."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(table["name"].upper() + " (")
    r.font.name = "Consolas"
    r.font.size = Pt(font)
    parts = formal(table)
    for i, (n, pk) in enumerate(parts):
        r = p.add_run(n)
        r.font.name = "Consolas"
        r.font.size = Pt(font)
        r.underline = pk
        if pk:
            r.bold = True
        if i < len(parts) - 1:
            r2 = p.add_run(", ")
            r2.font.name = "Consolas"
            r2.font.size = Pt(font)
    r = p.add_run(")")
    r.font.name = "Consolas"
    r.font.size = Pt(font)
    return p


def crop(src_pdf, rect_px, out_png, scale=2.0):
    d = pymupdf.open(src_pdf)
    x0, y0, x1, y1 = [v * 0.75 for v in rect_px]
    pix = d[0].get_pixmap(matrix=pymupdf.Matrix(scale, scale), clip=pymupdf.Rect(x0, y0, x1, y1))
    pix.save(out_png)


TB = {t["name"]: t for t in TABLES}
N_REL = len(TABLES)

# ================================================================= CAPA
for _ in range(4):
    P()
P("UNIVERSIDADE FEDERAL DE GOIÁS", bold=True, size=14, align="center", space_after=0)
P("Instituto de Informática – Engenharia de Software", size=12, align="center", space_after=0)
P("Disciplina: Banco de Dados 1", size=12, align="center", space_after=40)
P("PROJETO FINAL DE BANCO DE DADOS", bold=True, size=20, align="center", space_after=6)
P(SIS, bold=True, size=24, align="center", space_after=6)
P("Sistema Web de Gestão de Manutenção Industrial", size=13, align="center", space_after=30)
P("Entrega 1 – Projeto Lógico Conceitual", size=12, align="center", space_after=0)
P("Entrega 2 – Projeto Lógico Relacional", size=12, align="center", space_after=40)
P("Integrantes do grupo", bold=True, size=12, align="center", space_after=2)
for n in ["Ana Karolina da Silva Reges", "Matheus Vaz Teixeira", "Micael Henrique da Silva Fontes", "Reginaldo Ribeiro"]:
    P(n, size=12, align="center", space_after=0)
P()
P("Goiânia – setembro de 2026", size=12, align="center")
doc.add_page_break()

# ================================================================= SUMARIO
H("Sumário", 1)
p = doc.add_paragraph()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
r = OxmlElement("w:r")
t_ = OxmlElement("w:t")
t_.text = "Clique com o botão direito e escolha “Atualizar campo” para gerar o sumário."
r.append(t_)
fld.append(r)
p._p.append(fld)
doc.add_page_break()

# ================================================================= PARTE 1
H("PARTE 1 – PROJETO LÓGICO CONCEITUAL", 1)

H("1.1 Definição do objetivo", 2)
PJ(f"O objetivo deste projeto é modelar e implementar o banco de dados do sistema {SIS}, uma "
   "plataforma web que centraliza, digitaliza e padroniza a engenharia e a operação de manutenção de uma planta "
   "industrial. O banco deve sustentar o ciclo de vida completo das intervenções técnicas – abertura, distribuição, "
   "execução, aprovação e consolidação no histórico –, o controle do inventário interno de peças de reposição, o "
   "agendamento automático de manutenções preventivas e o cálculo de indicadores de confiabilidade (MTBF e MTTR).")
PJ("Do ponto de vista da disciplina, o trabalho percorre as três fases clássicas do projeto de banco de dados "
   "(Elmasri & Navathe): o projeto conceitual, expresso em um Modelo Entidade-Relacionamento Estendido; o projeto "
   "lógico relacional, obtido pelo mapeamento MER → Modelo Relacional e normalizado até a 3ª Forma Normal; e o "
   "projeto físico em PostgreSQL, previsto para a entrega final.")
P("Objetivos específicos do banco de dados", bold=True)
B("manter um histórico rastreável de todas as ordens de serviço executadas em cada equipamento;", "Rastreabilidade: ")
B("garantir que toda baixa de peça seja registrada e que o saldo em estoque nunca fique negativo;", "Integridade do estoque: ")
B("registrar a decisão de aprovação ou rejeição do supervisor, com justificativa obrigatória na rejeição;", "Fluxo de aprovação: ")
B("armazenar os dados brutos (datas de falha, início e fim de execução) que permitem calcular MTBF, MTTR e disponibilidade;", "Indicadores: ")
B("armazenar os planos de recorrência (periodicidade) que alimentam o agendamento automático de preventivas;", "Planejamento: ")
B("guardar somente os dados pessoais necessários, com senha armazenada apenas como hash, em conformidade com a LGPD.", "Segurança e LGPD: ")

H("1.2 Descrição do cenário e do escopo", 2)
H("1.2.1 Cenário", 3)
PJ("A área de negócio é a Gestão de Manutenção Industrial e de Ativos Físicos em plantas de manufatura com processos "
   "produtivos automatizados. Nesse ambiente, a disponibilidade contínua das máquinas determina a produtividade da "
   "fábrica. O diagnóstico realizado na fase de requisitos identificou cinco problemas centrais: paradas não planejadas "
   "por ausência de preventivas sistemáticas; burocracia de formulários em papel e planilhas dispersas; perda e "
   "inconsistência de dados nas trocas de turno; descontrole do almoxarifado interno de peças; e ausência de "
   "indicadores gerenciais para decidir sobre reforma ou substituição de equipamentos.")
PJ("O sistema atende quatro perfis de negócio, além de um perfil técnico de Administrador do Sistema, cada um com "
   "uma visão distinta dos mesmos dados:")
T(["Perfil", "Atuação no sistema", "Dados que produz / consome"], [
    ["Técnico de Manutenção", "Registra corretivas em campo, executa checklists de preventivas e dá baixa em peças.",
     "Produz: ordens de serviço, resultados de checklist, baixas de peça. Consome: histórico do equipamento, tarefas da semana."],
    ["Supervisor de Manutenção", "Distribui OS por prioridade e prazo, monitora o status da equipe, aprova ou rejeita OS executadas.",
     "Produz: OS distribuídas, decisões de aprovação/rejeição, equipes. Consome: calendário, status dos técnicos, consumo de peças."],
    ["Gestor Industrial", "Acompanha dashboards de MTBF/MTTR, ranking de máquinas críticas e produtividade; consulta relatórios executivos.",
     "Consome: todas as OS consolidadas, falhas por equipamento, tarefas planejadas × concluídas."],
    ["Equipe Administrativa", "Mantém cadastros de equipamentos, checklists e peças; acompanha alertas de estoque.",
     "Produz: equipamentos, itens de checklist, peças, linhas de produção. Consome: alertas de estoque mínimo."],
    ["Administrador do Sistema", "Superusuário de TI: gerencia contas de usuário e corrige dados em caso de problema.",
     "Produz: usuários. Consome: tudo (acesso total)."],
], widths=[3.5, 6.0, 6.5])
H("1.2.2 Escopo funcional coberto pelo banco de dados", 3)
for txt in ["Gestão de Ativos: cadastro de equipamentos (com fabricante, modelo, linha de produção e criticidade) e do checklist de preventiva de cada máquina.",
            "Ordens de Serviço: registro de manutenções corretivas e preventivas, com técnico responsável, prioridade, prazo, status e execução.",
            "Agendamento Preventivo: planos de recorrência por periodicidade (Diária, Semanal, Mensal, Trimestral, Semestral ou Anual), que geram OS preventivas.",
            "Auditoria e Fluxo de Trabalho: distribuição de OS pelo supervisor e aprovação/rejeição com justificativa antes da consolidação no histórico.",
            "Estoque Interno: peças de reposição com fornecedor homologado, estoque mínimo, saldo atual e baixa automática por OS.",
            "Indicadores e Relatórios: dados de falha e de execução necessários a MTBF, MTTR, disponibilidade, ranking, produtividade e horas trabalhadas.",
            "Notificações: alertas de proximidade/atraso de preventiva e de estoque mínimo, com registro de envio e leitura por usuário."]:
    B(txt)
H("1.2.3 Fora do escopo", 3)
for txt in ["Importação de dados por planilhas (CSV/Excel) e exportação de relatórios em PDF ou planilha: os relatórios são consultados na própria interface.",
            "Anexo de arquivos (fotos, documentos ou manuais) a ordens de serviço ou equipamentos.",
            "Compras, cotações e seleção de fornecedores (o fornecedor é apenas um dado descritivo da peça).",
            "Gestão financeira: custos monetários das OS, preço de peças, faturamento e notas fiscais (delegados ao ERP).",
            "Folha de pagamento e processamento financeiro de horas trabalhadas.",
            "Telemetria IoT: o horímetro é informado manualmente pelo técnico a cada manutenção.",
            "Inteligência artificial para previsão de falhas e aplicativo móvel nativo."]:
    B(txt)

H("1.3 Definição dos requisitos", 2)
PJ("Os requisitos abaixo reproduzem a Especificação de Requisitos do grupo (RF01–RF16, RNF01–RNF10), revisada em "
   "11/09/2026, com foco nos dados que cada requisito exige do banco.")
H("1.3.1 Requisitos funcionais", 3)
RF = [
    ("RF01", "Cadastro de Equipamentos", "nome, modelo, número de série (único), localização na fábrica, linha de produção e criticidade (Alta/Média/Baixa)."),
    ("RF02", "Registrar Manutenção Corretiva", "máquina (por código ou nome), descrição da falha (até 500 caracteres), causa, parada de produção, horímetro, peças trocadas com quantidade, data/hora automática."),
    ("RF03", "Registrar Manutenção Preventiva", "checklist específico da máquina, resultado por item (Verificado/Não aplicável), horímetro, observações, data/hora automática; conclusão bloqueada com item obrigatório pendente."),
    ("RF04", "Agendamento Automatizado", "plano por máquina com periodicidade Diária, Semanal, Mensal, Trimestral, Semestral ou Anual; próxima data calculada pela periodicidade; horímetro registrado apenas para histórico."),
    ("RF05", "Painel de Indicadores", "MTBF, MTTR e disponibilidade por máquina, calculados a partir das datas de falha, início e fim de execução das OS; atualização a cada 60 s."),
    ("RF06", "Alertas de Manutenção", "notificação a técnico e supervisor quando preventiva estiver a menos de 24 h do prazo ou atrasada há mais de 24 h."),
    ("RF07", "Geração de Relatórios", "histórico de falhas por equipamento e consumo de peças (quantidades) por período e por máquina, consultados em tela, sem custos."),
    ("RF08", "Consultar Histórico da Máquina", "todas as intervenções de um equipamento, por código ou nome parcial, da mais recente para a mais antiga."),
    ("RF09", "Visualizar Tarefas da Semana", "OS preventivas atribuídas ao técnico, ordenadas por prazo."),
    ("RF10", "Dar Baixa em Peças", "peça, quantidade, OS de origem; redução automática do saldo; saldo nunca negativo."),
    ("RF11", "Distribuir Ordens de Serviço", "supervisor define máquina, tipo, descrição, prioridade, prazo e técnico responsável."),
    ("RF12", "Aprovar Ordens de Serviço", "decisão Aprovada/Rejeitada, data, supervisor e justificativa obrigatória na rejeição."),
    ("RF13", "Alerta de Estoque Mínimo", "após cada baixa (RF10), se saldo ≤ estoque mínimo, notificar Administrativo e Supervisor; sem novo alerta para a mesma peça em 24 h."),
    ("RF14", "Monitorar Status da Equipe", "status Ocupado/Disponível de cada técnico, atualizado ao iniciar/finalizar OS e alterável pelo supervisor."),
    ("RF15", "Identificar Máquinas Críticas", "ranking das 10 máquinas com mais falhas no período (30/90/180 dias) com MTBF."),
    ("RF16", "Monitorar Produtividade", "tarefas planejadas × concluídas por período, equipe e técnico."),
]
T(["ID", "Requisito", "Dados exigidos do banco"], [list(r) for r in RF], widths=[1.3, 4.2, 10.5])
H("1.3.2 Requisitos não funcionais com impacto no banco de dados", 3)
T(["ID", "Requisito", "Impacto no projeto do banco"], [
    ["RNF03/RNF08", "Autenticação com login/senha; 2FA obrigatório para Supervisor, Gestor e Administrador do Sistema", "USUARIO guarda login único, hash da senha e flag de 2FA."],
    ["RNF09", "Conformidade com a LGPD", "Dados pessoais mínimos (nome, login, e-mail, telefone); exclusão/retificação suportadas."],
    ["RNF10", "Controle de acesso por perfil (RBAC) com os quatro perfis de negócio e o Administrador do Sistema", "Especialização de USUARIO em 5 subclasses disjuntas; discriminador tipo_usuario."],
    ["RNF04/RNF05", "Páginas em até 3 s; disponibilidade 99,5% no horário de produção", "Índices nas chaves estrangeiras e nas colunas de busca (código, número de série, datas)."],
    ["RNF02", "Registro de falha em menos de 5 cliques", "Campos da OS corretiva com valores padrão (data/hora automática, prioridade, status)."],
], widths=[2.3, 6.0, 7.7])
H("1.3.3 Regras de negócio", 3)
RN = [
    ("RN01", "Todo usuário tem exatamente um perfil: Técnico, Supervisor, Gestor, Administrativo ou Administrador do Sistema (especialização disjunta e total)."),
    ("RN02", "Toda ordem de serviço é corretiva ou preventiva, nunca ambas, e sempre pertence a um único equipamento."),
    ("RN03", "Toda OS tem um único técnico responsável; o supervisor que a distribuiu é opcional (corretivas podem ser abertas em campo pelo técnico)."),
    ("RN04", "Uma OS só é consolidada no histórico após decisão Aprovada do supervisor; a rejeição exige justificativa e devolve a OS ao técnico. A OS guarda a decisão vigente."),
    ("RN05", "Cada máquina tem seu próprio checklist de preventiva; todo item obrigatório precisa de resultado para concluir a OS."),
    ("RN06", "Um equipamento pode ter vários planos preventivos; cada plano tem uma periodicidade fixa e uma próxima data."),
    ("RN07", "Toda baixa de peça está ligada a uma OS e a uma peça, com quantidade positiva; o saldo da peça nunca fica negativo."),
    ("RN08", "Quando o saldo de uma peça fica igual ou abaixo do estoque mínimo, gera-se um alerta para Administrativo e Supervisor, no máximo um por peça a cada 24 h."),
    ("RN09", "Cada peça registra o nome de um fornecedor homologado; o processo de compra é externo ao sistema."),
    ("RN10", "Um técnico pertence a exatamente uma equipe; uma equipe é liderada por exatamente um supervisor."),
    ("RN11", "O status do técnico (Disponível/Ocupado) é atualizado automaticamente ao iniciar/encerrar uma OS, mas pode ser sobrescrito pelo supervisor; por isso é armazenado, não derivado."),
    ("RN12", "Códigos e números de série de equipamentos, códigos de peças, logins e e-mails são únicos no sistema."),
    ("RN13", "Não há valores monetários no banco (custos, preços e faturamento estão fora do escopo)."),
    ("RN14", "O horímetro é registrado em cada OS apenas como histórico; o agendamento das preventivas depende só da periodicidade do plano."),
]
T(["ID", "Regra"], [list(r) for r in RN], widths=[1.5, 14.5])

# ---------------------------------------------------------------- 1.4 prototipos
H("1.4 Protótipos para análise dos requisitos de dados", 2)
PJ("Os protótipos de baixa fidelidade abaixo foram usados para levantar os dados capturados e exibidos em cada tela. "
   "Cada campo do protótipo foi rastreado até um atributo ou relacionamento do modelo conceitual, o que garante que o "
   "DER cobre todas as informações que o sistema precisa persistir.")
PROTOS = [
    ("P01", "Cadastro de Equipamento e do seu checklist (Administrativo – RF01, RF03)", "P01_cadastro_equipamento.png",
     "EQUIPAMENTO (código, nome, fabricante, modelo, número de série, localização, criticidade, horímetro, data de aquisição); fica_na_linha → LINHA_PRODUCAO; ITEM_CHECKLIST (entidade fraca: número, descrição, obrigatório)."),
    ("P02", "Registro de Manutenção Corretiva (Técnico – RF02, RF10)", "P02_registrar_corretiva.png",
     "ORDEM_SERVICO (descrição, data de abertura, prioridade, leitura do horímetro) + MANUTENCAO_CORRETIVA (causa, impacto, data/hora da falha); consome → PECA (quantidade, data da baixa); sofre → EQUIPAMENTO; executa → TECNICO."),
    ("P03", "Execução de Checklist de Preventiva (Técnico – RF03, RF09)", "P03_checklist_preventiva.png",
     "MANUTENCAO_PREVENTIVA (observação geral) gerada por PLANO_PREVENTIVO; verifica → ITEM_CHECKLIST do equipamento (resultado, observação); datas de início/fim de execução e horímetro da OS."),
    ("P04", "Aprovação de Ordem de Serviço (Supervisor – RF12)", "P04_aprovar_os.png",
     "Relacionamento avalia (SUPERVISOR – ORDEM_SERVICO) com decisão, data e justificativa; mudança de status da OS."),
    ("P05", "Cadastro de Peça e Alertas de Estoque (Administrativo – RF10, RF13)", "P05_cadastro_peca.png",
     "PECA (código, descrição, unidade, fornecedor, estoque mínimo, quantidade atual); NOTIFICACAO do tipo Estoque mínimo com refere_peca."),
    ("P06", "Dashboard de Indicadores (Gestor – RF05, RF15, RF16)", "P06_dashboard.png",
     "Somente leitura: agrega ORDEM_SERVICO/MANUTENCAO_CORRETIVA (datas de falha e execução → MTBF, MTTR), EQUIPAMENTO, LINHA_PRODUCAO, EQUIPE e TECNICO (planejadas × concluídas)."),
]
for pid, titulo, arq, dados in PROTOS:
    H(f"{pid} – {titulo}", 3)
    IMG(f"diagramas/prototipos/{arq}", 16.0)
    P("Dados identificados: ", bold=True, space_after=2).add_run(dados).bold = False
T(["Protótipo", "Entidades alimentadas", "Relacionamentos exercitados"], [
    ["P01", "EQUIPAMENTO, ITEM_CHECKLIST, LINHA_PRODUCAO", "fica_na_linha, possui_item"],
    ["P02", "ORDEM_SERVICO, MANUTENCAO_CORRETIVA, PECA", "sofre, executa, consome"],
    ["P03", "MANUTENCAO_PREVENTIVA, ITEM_CHECKLIST, PLANO_PREVENTIVO", "gera, verifica, possui_item"],
    ["P04", "ORDEM_SERVICO", "avalia"],
    ["P05", "PECA, NOTIFICACAO", "refere_peca, recebe"],
    ["P06", "(leitura) ORDEM_SERVICO, EQUIPAMENTO, EQUIPE, TECNICO", "sofre, executa, pertence, lidera"],
], widths=[2.0, 7.0, 7.0])

# ---------------------------------------------------------------- 1.5 entidades
H("1.5 Descrição das entidades e seus atributos", 2)
PJ("Convenções: (PK) atributo-chave; (PP) chave parcial de entidade fraca; (D) atributo derivado; (E) atributo de "
   "domínio enumerado. Os tipos são conceituais; os tipos físicos aparecem no dicionário de dados da Parte 2. "
   "Chaves estrangeiras e atributos de relacionamentos não aparecem aqui: pertencem ao modelo relacional e à tabela "
   "de relacionamentos, respectivamente.")

REL_ATTRS_OS = ("decisao_avaliacao", "data_avaliacao", "justificativa")


def ctype(pg):
    pg = pg.upper()
    if pg.startswith("VARCHAR") or pg.startswith("CHAR") or pg == "TEXT":
        return "texto"
    if pg in ("SERIAL", "INTEGER", "SMALLINT"):
        return "inteiro"
    if pg.startswith("NUMERIC"):
        return "numérico"
    if pg == "BOOLEAN":
        return "lógico"
    if pg == "DATE":
        return "data"
    if pg == "TIMESTAMP":
        return "data/hora"
    return pg.lower()


def ent_table(title, desc, table_name, skip=(), extra=(), pk_partial=False):
    H(title, 3)
    PJ(desc)
    rows = []
    for c in TB[table_name]["cols"]:
        if c[0] in skip or fk_target(c):
            continue
        tags = []
        if is_pk(c):
            tags.append("PP" if pk_partial else "PK")
        if any(f.startswith("CK:") and " IN (" in f for f in flags(c)):
            tags.append("E")
        rows.append([c[0], ctype(c[1]), c[3], ", ".join(tags)])
    for e in extra:
        rows.append(list(e))
    if rows:
        T(["Atributo", "Tipo", "Descrição", "Obs."], rows, widths=[4.2, 2.3, 8.3, 1.2])
    else:
        P("Sem atributos próprios: herda todos os atributos de USUARIO.", italic=True, size=10)


H("1.5.1 Módulo Pessoas", 3)
ent_table("USUARIO (superclasse)", "Qualquer pessoa autenticada no sistema. Superclasse dos cinco perfis; concentra os dados comuns de "
          "identificação e acesso. Especialização disjunta e total: todo usuário é membro de exatamente uma subclasse.",
          "usuario", skip=("tipo_usuario",))
ent_table("TECNICO (subclasse de USUARIO)", "Executor operacional das manutenções. Herda os atributos de USUARIO; participa de executa (OS) e pertence (EQUIPE).",
          "tecnico")
ent_table("SUPERVISOR (subclasse de USUARIO)", "Responsável pela distribuição de OS, pela equipe e pela qualidade dos registros. Participa de lidera, distribui e avalia.",
          "supervisor")
ent_table("GESTOR (subclasse de USUARIO)", "Perfil estratégico que consome indicadores; não altera dados operacionais.", "gestor")
ent_table("ADMINISTRATIVO (subclasse de USUARIO)", "Mantém os cadastros-base (equipamentos, checklists, peças).", "administrativo")
ent_table("ADMINISTRADOR_SISTEMA (subclasse de USUARIO)", "Perfil técnico de TI (superusuário) que gerencia contas e corrige dados. "
          "Existe como subclasse para que a especialização continue total e para o controle de acesso (RBAC).", "administrador_sistema")
ent_table("EQUIPE", "Grupo de técnicos liderado por um supervisor; usado no filtro de produtividade por equipe (RF16) e no painel de status (RF14).",
          "equipe")

H("1.5.2 Módulo Ativos", 3)
ent_table("LINHA_PRODUCAO", "Agrupamento produtivo dos equipamentos; filtro do dashboard (RF05).", "linha_producao")
ent_table("EQUIPAMENTO", "Máquina industrial sob manutenção, identificada por número de série único e por um código interno usado nas buscas. "
          "Fabricante e modelo são descritivos. Possui o histórico de OS (sofre) e o próprio checklist de preventiva (possui_item).",
          "equipamento")
ent_table("ITEM_CHECKLIST (entidade fraca de EQUIPAMENTO)", "Passo de inspeção do checklist de uma máquina. Não existe sem o equipamento; "
          "identificada pelo equipamento + número do item (relacionamento identificador possui_item).", "item_checklist", pk_partial=True)
ent_table("PLANO_PREVENTIVO", "Regra de recorrência de preventiva de um equipamento, com periodicidade fixa e próxima data. "
          "Cada vencimento gera uma MANUTENCAO_PREVENTIVA.", "plano_preventivo")

H("1.5.3 Módulo Manutenção", 3)
ent_table("ORDEM_SERVICO (superclasse)", "Registro formal de uma intervenção técnica em um equipamento. Especialização disjunta e total em "
          "corretiva e preventiva. Participa de sofre (EQUIPAMENTO), executa (TECNICO), distribui (SUPERVISOR), avalia (SUPERVISOR), "
          "consome (PECA) e refere_os (NOTIFICACAO).", "ordem_servico", skip=("tipo_os",) + REL_ATTRS_OS,
          extra=[("duracao_execucao", "numérico", "Horas trabalhadas = data_fim_execucao − data_inicio_execucao (base do MTTR e do relatório de horas)", "D")])
ent_table("MANUTENCAO_CORRETIVA (subclasse de ORDEM_SERVICO)", "OS aberta em resposta a uma falha; guarda os dados que alimentam o MTBF e o ranking de máquinas críticas.",
          "manutencao_corretiva")
ent_table("MANUTENCAO_PREVENTIVA (subclasse de ORDEM_SERVICO)", "OS programada, normalmente gerada por um plano preventivo; exige o checklist do equipamento (verifica).",
          "manutencao_preventiva")

H("1.5.4 Módulo Estoque", 3)
ent_table("PECA", "Insumo de reposição com controle de estoque mínimo e saldo atual; registra o nome do fornecedor homologado. Participa de consome e refere_peca.", "peca")

H("1.5.5 Módulo Notificação", 3)
ent_table("NOTIFICACAO", "Aviso gerado automaticamente pelo sistema (proximidade de prazo, atraso, estoque mínimo, nova OS, aprovação, rejeição). "
          "Pode referir-se a uma OS ou a uma peça; é entregue a um ou mais usuários (recebe). A data de geração permite suprimir alertas repetidos em 24 h.", "notificacao")

H("1.5.6 Relacionamentos", 3)
REL = [
    ("lidera", "SUPERVISOR – EQUIPE", "1:N", "Equipe: total", "–", "Um supervisor lidera várias equipes; toda equipe tem um supervisor."),
    ("pertence", "TECNICO – EQUIPE", "N:1", "Técnico: total", "–", "Todo técnico pertence a exatamente uma equipe."),
    ("fica_na_linha", "EQUIPAMENTO – LINHA_PRODUCAO", "N:1", "Equipamento: total", "–", "Todo equipamento está em uma linha de produção."),
    ("possui_item", "EQUIPAMENTO – ITEM_CHECKLIST", "1:N", "Item: total (identificador)", "–", "Relacionamento identificador da entidade fraca ITEM_CHECKLIST."),
    ("tem_plano", "EQUIPAMENTO – PLANO_PREVENTIVO", "1:N", "Plano: total", "–", "Um equipamento pode ter vários planos; todo plano é de um equipamento."),
    ("sofre", "EQUIPAMENTO – ORDEM_SERVICO", "1:N", "OS: total", "–", "Histórico de intervenções do equipamento."),
    ("executa", "TECNICO – ORDEM_SERVICO", "1:N", "OS: total", "–", "Toda OS tem um técnico responsável."),
    ("distribui", "SUPERVISOR – ORDEM_SERVICO", "1:N", "parcial", "–", "OS distribuída pelo supervisor (RF11); corretivas abertas em campo não têm distribuidor."),
    ("avalia", "SUPERVISOR – ORDEM_SERVICO", "1:N", "parcial", "decisao, data_avaliacao, justificativa", "Decisão vigente do supervisor sobre a OS (RF12); justificativa obrigatória na rejeição."),
    ("gera", "PLANO_PREVENTIVO – MANUTENCAO_PREVENTIVA", "1:N", "parcial", "–", "Preventiva gerada por um plano; preventivas avulsas não têm plano."),
    ("verifica", "MANUTENCAO_PREVENTIVA – ITEM_CHECKLIST", "N:M", "parcial", "resultado, observacao", "Resultado de cada item do checklist na execução da preventiva."),
    ("consome", "ORDEM_SERVICO – PECA", "N:M", "parcial", "quantidade, data_baixa", "Baixa de estoque: peças consumidas por uma OS."),
    ("refere_os", "NOTIFICACAO – ORDEM_SERVICO", "N:1", "parcial", "–", "Notificação originada por uma OS (proximidade, atraso, aprovação...)."),
    ("refere_peca", "NOTIFICACAO – PECA", "N:1", "parcial", "–", "Notificação originada por uma peça (estoque mínimo)."),
    ("recebe", "NOTIFICACAO – USUARIO", "N:M", "Notificação: total", "canal, data_envio, data_leitura", "Destinatários da notificação e registro de envio/leitura."),
]
T(["Relacionamento", "Entidades", "Card.", "Participação", "Atributos", "Descrição"], [list(r) for r in REL],
  widths=[2.5, 4.0, 1.0, 2.6, 2.4, 3.5], font=8)
H("1.5.7 Especializações (MER Estendido)", 3)
T(["Superclasse", "Subclasses", "Restrições", "Atributo definidor", "Justificativa"], [
    ["USUARIO", "TECNICO, SUPERVISOR, GESTOR, ADMINISTRATIVO, ADMINISTRADOR_SISTEMA", "disjunta (d), total", "tipo_usuario",
     "Cada perfil tem atributos e relacionamentos próprios (status do técnico, equipe do supervisor, avaliações) e permissões distintas (RBAC)."],
    ["ORDEM_SERVICO", "MANUTENCAO_CORRETIVA, MANUTENCAO_PREVENTIVA", "disjunta (d), total", "tipo_os",
     "Corretiva registra causa da falha e impacto; preventiva liga-se ao plano e ao checklist. Os atributos comuns (técnico, equipamento, datas, status, avaliação) ficam na superclasse."],
], widths=[2.5, 4.0, 2.2, 2.3, 5.0], font=8)

# ---------------------------------------------------------------- 1.6 DER
landscape(True)
H("1.6 Modelo Entidade-Relacionamento (DER)", 2)
PJ("Diagrama completo na notação de Chen adotada nas aulas (Elmasri & Navathe): retângulos = entidades; retângulo duplo = "
   "entidade fraca; losangos = relacionamentos; losango duplo = relacionamento identificador; elipses = atributos "
   "(sublinhado = chave; sublinhado tracejado = chave parcial; elipse tracejada = derivado); "
   "linha dupla = participação total; círculo (d) = especialização disjunta; 1, N, M = razão de cardinalidade.")
doc.add_page_break()
IMG("diagramas/01_der_conceitual.png", 24.5, f"Figura 1 – DER conceitual completo do {SIS}")
doc.add_page_break()
crop("ferramentas/_build/01_der_conceitual.pdf", (0, 40, 1010, 1130), "ferramentas/_build/der_pessoas.png")
crop("ferramentas/_build/01_der_conceitual.pdf", (380, 380, 2200, 1130), "ferramentas/_build/der_manutencao.png")
crop("ferramentas/_build/01_der_conceitual.pdf", (1180, 40, 2750, 1000), "ferramentas/_build/der_ativos.png")
crop("ferramentas/_build/01_der_conceitual.pdf", (450, 940, 2750, 1400), "ferramentas/_build/der_estoque.png")
H("Recortes por módulo (para leitura)", 3)
IMG("ferramentas/_build/der_pessoas.png", 12.5, "Figura 1a – Módulo Pessoas: USUARIO e sua especialização, EQUIPE, avalia")
doc.add_page_break()
IMG("ferramentas/_build/der_manutencao.png", 24.5, "Figura 1b – Módulo Manutenção: ORDEM_SERVICO, subclasses, avalia, consome")
doc.add_page_break()
IMG("ferramentas/_build/der_ativos.png", 22.0, "Figura 1c – Módulo Ativos: EQUIPAMENTO, ITEM_CHECKLIST, PLANO_PREVENTIVO, LINHA_PRODUCAO")
doc.add_page_break()
IMG("ferramentas/_build/der_estoque.png", 24.5, "Figura 1d – Módulos Estoque e Notificação: PECA, consome, NOTIFICACAO, recebe")
landscape(False)

# ================================================================= PARTE 2
H("PARTE 2 – PROJETO LÓGICO RELACIONAL", 1)
H("2.1 Modelo relacional normalizado (após o mapeamento)", 2)
PJ("O mapeamento seguiu o algoritmo de sete passos de Elmasri & Navathe apresentado na aula 05, acrescido do passo de "
   "especialização (aula 03). Notação formal adotada: NOME_DA_RELAÇÃO (atributos); a chave primária aparece sublinhada em "
   "negrito e as chaves estrangeiras são prefixadas com #, como no exercício de mapeamento resolvido em aula.")

H("2.1.1 Mapeamento passo a passo", 3)
P("Passo 1 – Entidades fortes → uma relação por entidade, com atributos simples e chave primária.", bold=True)
for n in ["usuario", "equipe", "linha_producao", "equipamento", "plano_preventivo", "ordem_servico", "peca", "notificacao"]:
    FORMAL(TB[n])
P("Obs.: as chaves estrangeiras exibidas com # são acrescentadas no passo 4; aparecem aqui já na forma final para "
  "evitar repetição. O atributo derivado duracao_execucao não é armazenado: é calculado a partir de data_inicio_execucao e data_fim_execucao.", italic=True, size=9)
P("Passo 2 – Entidades fracas → relação com os atributos da entidade fraca + chave primária da entidade proprietária (FK); "
  "a PK é a composição da chave do proprietário com a chave parcial.", bold=True)
FORMAL(TB["item_checklist"])
P("Passo 3 – Relacionamentos 1:1: não há relacionamentos binários 1:1 no modelo.", bold=True)
P("Passo 4 – Relacionamentos 1:N → a chave primária do lado 1 vai como chave estrangeira para a relação do lado N, "
  "junto com os atributos do relacionamento, se houver. A FK é obrigatória (NOT NULL) quando a participação do lado N é total "
  "e opcional quando é parcial.", bold=True)
T(["Relacionamento", "Lado 1 → lado N", "Acrescentado em", "Nula?"], [
    ["lidera", "SUPERVISOR → EQUIPE", "equipe.id_supervisor", "não"],
    ["pertence", "EQUIPE → TECNICO", "tecnico.id_equipe", "não"],
    ["fica_na_linha", "LINHA_PRODUCAO → EQUIPAMENTO", "equipamento.id_linha", "não"],
    ["tem_plano", "EQUIPAMENTO → PLANO_PREVENTIVO", "plano_preventivo.id_equipamento", "não"],
    ["sofre", "EQUIPAMENTO → ORDEM_SERVICO", "ordem_servico.id_equipamento", "não"],
    ["executa", "TECNICO → ORDEM_SERVICO", "ordem_servico.id_tecnico", "não"],
    ["distribui", "SUPERVISOR → ORDEM_SERVICO", "ordem_servico.id_supervisor", "sim (participação parcial)"],
    ["avalia", "SUPERVISOR → ORDEM_SERVICO", "ordem_servico.id_supervisor_avaliador + atributos decisao_avaliacao, data_avaliacao, justificativa", "sim (participação parcial)"],
    ["gera", "PLANO_PREVENTIVO → MANUTENCAO_PREVENTIVA", "manutencao_preventiva.id_plano", "sim (participação parcial)"],
    ["refere_os", "ORDEM_SERVICO → NOTIFICACAO", "notificacao.codigo_os", "sim (participação parcial)"],
    ["refere_peca", "PECA → NOTIFICACAO", "notificacao.id_peca", "sim (participação parcial)"],
], widths=[2.6, 5.4, 6.0, 2.0])
P("Passo 5 – Relacionamentos N:M → nova relação com as chaves primárias das duas entidades (FKs) e os atributos do relacionamento; "
  "a PK é a composição das duas FKs.", bold=True)
for n in ["baixa_peca", "resultado_checklist", "notificacao_usuario"]:
    FORMAL(TB[n])
P("Passo 6 – Atributos multivalorados: não há no modelo (o telefone do usuário é monovalorado).", bold=True)
P("Passo 7 – Relacionamentos n-ários (grau > 2): não há no modelo.", bold=True)
P("Passo 8 – Especialização/generalização (opção 8A da aula 03: uma relação para a superclasse e uma para cada subclasse, "
  "compartilhando a chave primária, mais um atributo discriminador na superclasse).", bold=True)
for n in ["tecnico", "supervisor", "gestor", "administrativo", "administrador_sistema", "manutencao_corretiva", "manutencao_preventiva"]:
    FORMAL(TB[n])
P("A opção 8A foi escolhida porque as duas especializações são totais e disjuntas, as subclasses têm atributos e "
  "relacionamentos próprios, e ela evita colunas nulas (opção 8C) sem duplicar os atributos comuns (opção 8B). "
  "ADMINISTRADOR_SISTEMA fica só com a chave, o que é suficiente para o controle de acesso.", italic=True, size=9)

H("2.1.2 Esquema relacional final (3FN)", 3)
PJ(f"Esquema S = {{R1, ..., R{N_REL}}} resultante do mapeamento, já com as chaves estrangeiras incorporadas:")
for mod in ["Pessoas", "Ativos", "Manutencao", "Estoque", "Notificacao"]:
    P(f"Módulo {mod}", bold=True, space_after=2)
    for t in TABLES:
        if t["modulo"] == mod:
            FORMAL(t)

H("2.1.3 Verificação da normalização", 3)
PJ("Exemplo de decomposição: a ficha de manutenção em papel usada hoje na fábrica corresponde a uma relação não normalizada. "
   "O mesmo raciocínio das aulas 06 (BOLETIM e PEDIDO) foi aplicado:")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(0.5)
r = p.add_run("FICHA_OS (codigo_os, descricao, data_abertura, codigo_equip, nome_equip, numero_serie, linha, "
              "login_tecnico, nome_tecnico, equipe, {codigo_peca, descricao_peca, fornecedor, quantidade})")
r.font.name = "Consolas"
r.font.size = Pt(9.5)
B("o grupo repetitivo {peça} viola a 1FN → extraído para BAIXA_PECA (codigo_os, id_peca, quantidade, data_baixa), com PK composta.", "1FN: ")
B("em BAIXA_PECA, descricao_peca e fornecedor dependem só de id_peca (parte da chave) → movidos para PECA; em FICHA_OS não há dependência parcial porque a chave é simples.", "2FN: ")
B("nome_equip, numero_serie e linha dependem de codigo_equip, não de codigo_os (dependência transitiva) → EQUIPAMENTO e LINHA_PRODUCAO; nome_tecnico e equipe dependem de login_tecnico → USUARIO/TECNICO/EQUIPE.", "3FN: ")
PJ("Aplicando o mesmo critério a todas as relações do esquema final:")
NORM = [
    ["usuario", "id_usuario", "id_usuario → todos; login → id_usuario; email → id_usuario (chaves candidatas)", "Sim: chave simples; nenhum não-chave determina outro não-chave."],
    ["equipe", "id_equipe", "id_equipe → nome, turno, id_supervisor", "Sim."],
    ["tecnico / supervisor / gestor / administrativo / administrador_sistema", "id_usuario", "id_usuario → atributos da subclasse", "Sim."],
    ["linha_producao", "id_linha", "id_linha → nome, setor", "Sim."],
    ["equipamento", "id_equipamento", "id_equipamento → todos; numero_serie → id_equipamento; codigo → id_equipamento", "Sim: fabricante e modelo dependem só do equipamento; nome da linha não é repetido (transitivo via id_linha)."],
    ["item_checklist", "id_equipamento, num_item", "(id_equipamento, num_item) → descricao, obrigatorio", "Sim: descricao/obrigatorio dependem da chave inteira (2FN); sem transitividade."],
    ["plano_preventivo", "id_plano", "id_plano → todos", "Sim."],
    ["ordem_servico", "codigo_os", "codigo_os → todos", "Sim: nome do técnico e do equipamento não são armazenados (transitivos); os atributos da avaliação dependem só da OS."],
    ["manutencao_corretiva / manutencao_preventiva", "codigo_os", "codigo_os → atributos da subclasse", "Sim."],
    ["resultado_checklist", "codigo_os, id_equipamento, num_item", "chave → resultado, observacao", "Sim: descricao do item não é repetida (dependeria só de parte da chave)."],
    ["peca", "id_peca", "id_peca → todos; codigo → id_peca", "Sim: fornecedor é um atributo simples da peça (nome), sem atributos próprios que gerassem transitividade."],
    ["baixa_peca", "codigo_os, id_peca", "chave → quantidade, data_baixa", "Sim: nenhum atributo depende de parte da chave."],
    ["notificacao", "id_notificacao", "id_notificacao → tipo, mensagem, data_geracao, codigo_os, id_peca", "Sim."],
    ["notificacao_usuario", "id_notificacao, id_usuario", "chave → canal, data_envio, data_leitura", "Sim."],
]
T(["Relação", "Chave primária", "Dependências funcionais", "3FN?"], NORM, widths=[3.6, 2.8, 5.6, 4.0], font=8)
PJ("Todas as relações estão na 1FN (atributos atômicos, sem grupos repetitivos), na 2FN (nas chaves compostas, nenhum "
   "atributo não-chave depende de parte da chave) e na 3FN (nenhum atributo não-chave depende transitivamente da chave). "
   "As dependências funcionais de cada relação partem apenas de chaves candidatas, o que também satisfaz a FNBC.")

# ---------------------------------------------------------------- 2.2 DER relacional
landscape(True)
H("2.2 Modelo relacional de forma gráfica (DER relacional)", 2)
PJ("Diagrama do esquema lógico em notação Engenharia da Informação (pé-de-galinha): cada caixa é uma relação com sua "
   "chave primária destacada; as linhas representam as chaves estrangeiras, com o lado “1” (traço) na relação referenciada "
   "e o lado “N” (pé-de-galinha) na relação que contém a FK; círculo = participação opcional (FK que aceita nulo). "
   "As cores agrupam as relações por módulo.")
doc.add_page_break()
IMG("diagramas/02_der_relacional.png", 16.0, f"Figura 2 – DER relacional (modelo lógico) do {SIS}")
landscape(False)

# ---------------------------------------------------------------- 2.3 dicionario
H("2.3 Dicionário de dados", 2)
PJ("Tipos de dados no padrão PostgreSQL (SGBD escolhido para o projeto físico). Convenções: PK = chave primária; "
   "FK = chave estrangeira (com a relação/coluna referenciada); UK = valor único; CK = restrição de domínio (CHECK); "
   "DEF = valor padrão. “Nulo = N” significa NOT NULL.")


def keydesc(c):
    parts = []
    for f in flags(c):
        if f == "PK":
            parts.append("PK")
        elif f.startswith("FK:") or f.startswith("FK2:"):
            parts.append("FK → " + f.split(":", 1)[1])
        elif f.startswith("UK"):
            parts.append("UK" + (" (composta)" if f == "UK2" else ""))
        elif f.startswith("CK:"):
            parts.append("CK: " + f[3:])
        elif f.startswith("DEF:"):
            parts.append("DEF " + f[4:])
    return "; ".join(parts)


for i, t in enumerate(TABLES, 1):
    H(f"2.3.{i} {t['name']}", 3)
    P("Origem no modelo conceitual: ", bold=True, space_after=2).add_run(t["origem"]).bold = False
    rows = [[c[0], c[1], "S" if c[2] else "N", keydesc(c), c[3]] for c in t["cols"]]
    T(["Atributo", "Tipo", "Nulo", "Chaves / restrições", "Descrição"], rows, widths=[3.4, 2.6, 0.9, 4.6, 4.5], font=8)
    fk_list = fks(t)
    if fk_list:
        P("Integridade referencial: " + "; ".join(
            f"({', '.join(l)}) referencia {rt}({', '.join(rc)})" for l, rt, rc in fk_list) + ".", size=9, italic=True)

# ---------------------------------------------------------------- referencias
H("Referências", 1)
for ref in ["ELMASRI, R.; NAVATHE, S. B. Sistemas de Banco de Dados. 6. ed. São Paulo: Pearson Addison-Wesley, 2011.",
            "LOUREIRO, A. C. B. Notas de aula de Banco de Dados 1: MER Estendido; Modelo Relacional; Mapeamento MER–MR; Normalização. UFG/INF, 2026.",
            "GRUPO GERIR MANUTENÇÃO. Descrição do Negócio, Escopo do Projeto e Especificação de Requisitos de Software (ERS) – revisão de setembro de 2026. UFG/INF, 2026.",
            "POSTGRESQL GLOBAL DEVELOPMENT GROUP. PostgreSQL 16 Documentation. Disponível em: https://www.postgresql.org/docs/16/."]:
    B(ref)

doc.save(OUT)
print("saved", OUT, "| relacoes:", N_REL)
